"""python-docx によるWordテキスト抽出。"""

import io

import docx

from .types import OcrResult, PageResult


def build_docx_block_index(file_bytes: bytes, document_id: str) -> list[dict]:
    """Build stable paragraph/table anchors matching the current preview order."""
    doc = docx.Document(io.BytesIO(file_bytes))
    blocks: list[dict] = []
    paragraph_index = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        blocks.append({
            "document_id": document_id,
            "type": "docx_block",
            "block_kind": "paragraph",
            "paragraph_index": paragraph_index,
            "text": text,
            "anchor_id": f"p-{paragraph_index}",
        })
        paragraph_index += 1

    table_index = 0
    for table in doc.tables:
        for row_index, row in enumerate(table.rows):
            for col_index, cell in enumerate(row.cells):
                text = cell.text.strip()
                if not text:
                    continue
                blocks.append({
                    "document_id": document_id,
                    "type": "docx_block",
                    "block_kind": "table_cell",
                    "table_index": table_index,
                    "row": row_index,
                    "col": col_index,
                    "text": text,
                    "anchor_id": f"t-{table_index}-r-{row_index}-c-{col_index}",
                })
        table_index += 1
    return blocks


def extract_docx(file_bytes: bytes, document_id: str) -> OcrResult:
    """Wordからテキストを抽出。文書全体=1ページ。"""
    doc = docx.Document(io.BytesIO(file_bytes))
    parts = []
    # 段落
    paragraph_index = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(f"[p-{paragraph_index}] {text}")
            paragraph_index += 1
    # テーブル
    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            cells = [
                f"[t-{table_index}-r-{row_index}-c-{col_index}] {cell.text.strip()}"
                for col_index, cell in enumerate(row.cells)
                if cell.text.strip()
            ]
            if cells:
                parts.append("\t".join(cells))
    text = "\n".join(parts)
    pages = [PageResult(page_number=1, text=text)] if text.strip() else []
    return OcrResult(document_id=document_id, pages=pages)
