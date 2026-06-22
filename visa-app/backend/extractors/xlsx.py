"""openpyxl によるExcelテキスト抽出。"""

import io

import openpyxl
from openpyxl.utils import get_column_letter

from .types import OcrResult, PageResult


def build_xlsx_cell_index(file_bytes: bytes, document_id: str) -> list[dict]:
    """Build stable cell anchors for resolver and HTML preview."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    cells: list[dict] = []
    for sheet_index, ws in enumerate(wb.worksheets, start=1):
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value)
                cell_ref = f"{get_column_letter(cell.column)}{cell.row}"
                cells.append({
                    "document_id": document_id,
                    "type": "xlsx_cell",
                    "sheet_name": ws.title,
                    "sheet_index": sheet_index,
                    "cell": cell_ref,
                    "row": cell.row,
                    "col": cell.column,
                    "text": text,
                    "anchor_id": f"{ws.title}!{cell_ref}",
                })
    wb.close()
    return cells


def extract_xlsx(file_bytes: bytes, document_id: str, sheet_name: str | None = None) -> OcrResult:
    """Excelからテキストを抽出。シート=ページとして扱う。"""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheets = [wb[sheet_name]] if sheet_name and sheet_name in wb.sheetnames else wb.worksheets
    pages = []
    for i, ws in enumerate(sheets):
        lines = []
        for row in ws.iter_rows(values_only=False):
            cells = []
            for cell in row:
                if cell.value is not None:
                    cells.append(str(cell.value))
            if cells:
                lines.append("\t".join(cells))
        text = "\n".join(lines)
        if text.strip():
            pages.append(PageResult(page_number=i + 1, text=f"[Sheet: {ws.title}]\n{text}"))
    wb.close()
    return OcrResult(document_id=document_id, pages=pages)
