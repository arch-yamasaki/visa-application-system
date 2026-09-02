#!/usr/bin/env python3
"""Build one self-contained HTML evidence report for targeted eval mismatches."""

from __future__ import annotations

import argparse
import base64
import html
import io
import json
import re
import sys
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]

TARGET_PATHS = {
    "applicant.birth_place",
    "applicant.home_country_address",
    "employer.address",
    "employer.corporate_number",
    "employer.has_corporate_number",
    "employment.has_position",
    "employment.position_title",
}

ADDRESS_PATHS = {
    "applicant.birth_place",
    "applicant.home_country_address",
    "employer.address",
}

CORPORATE_PATHS = {
    "employer.corporate_number",
    "employer.has_corporate_number",
}

POSITION_PATHS = {
    "employment.has_position",
    "employment.position_title",
}

FIELD_LABELS = {
    "applicant.birth_place": "出生地",
    "applicant.home_country_address": "本国住所",
    "employer.address": "所属機関住所",
    "employer.corporate_number": "法人番号",
    "employer.has_corporate_number": "法人番号の有無",
    "employment.has_position": "役職の有無",
    "employment.position_title": "役職名",
}

RASENS_PAGE_HINTS = {
    "applicant.birth_place": 1,
    "applicant.home_country_address": 1,
    "employer.address": 6,
    "employer.corporate_number": 6,
    "employer.has_corporate_number": 6,
    "employment.has_position": 6,
    "employment.position_title": 6,
}


@dataclass
class EvidenceImage:
    title: str
    badge: str
    image_uri: str
    caption: str


@dataclass
class FieldMismatch:
    path: str
    expected: str
    generated: str
    category: str


@dataclass
class Card:
    fixture_label: str
    case_dir: Path
    generated_dir: Path
    fields: list[FieldMismatch]
    category: str
    issue_title: str
    cause: str
    improvement: str
    rasens: EvidenceImage
    ai_source: EvidenceImage
    correct_source: EvidenceImage


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def normalize_text(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = re.sub(r"\s+", "", text).casefold()
    return text


def strip_md_escapes(value: str) -> str:
    return value.replace("\\|", "|").replace("<br>", "\n").strip()


def parse_detail_rows(report_path: Path) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    in_detail = False
    header: list[str] | None = None
    for line in report_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("### 全項目詳細"):
            in_detail = True
            continue
        if not in_detail:
            continue
        if line.startswith("| 大項目"):
            header = [c.strip() for c in line.strip("|").split("|")]
            continue
        if line.startswith("|---"):
            continue
        if line.startswith("---"):
            break
        if line.startswith("|") and header:
            cells = [strip_md_escapes(c.strip()) for c in line.strip("|").split("|")]
            if len(cells) >= len(header):
                rows.append(dict(zip(header, cells)))
    return rows


def find_fixture_dir(case_id: str) -> Path:
    for scenario in (ROOT / "visa-eval/test_cases_from_raw").glob("*/*/scenario.json"):
        try:
            if load_json(scenario).get("case_id") == case_id:
                return scenario.parent
        except Exception:
            continue
    raise FileNotFoundError(f"fixture not found for {case_id}")


def build_document_map(case_dir: Path) -> dict[str, dict[str, Any]]:
    docs: dict[str, dict[str, Any]] = {}
    for manifest_name, key in [
        ("input/document_manifest.json", "documents"),
        ("output/output_manifest.json", "outputs"),
    ]:
        manifest_path = case_dir / manifest_name
        if not manifest_path.exists():
            continue
        manifest = load_json(manifest_path)
        for doc in manifest.get(key, []):
            doc = dict(doc)
            doc["abs_path"] = (ROOT / doc["path"]).resolve()
            docs[doc["document_id"]] = doc
    return docs


def get_case_value(data: Any, path: str) -> Any:
    current = data
    for part in re.split(r"\.(?![^\[]*\])", path):
        m = re.fullmatch(r"([^\[]+)\[(\d+)\]", part)
        if m:
            key, idx = m.group(1), int(m.group(2))
            current = current.get(key, []) if isinstance(current, dict) else []
            current = current[idx] if isinstance(current, list) and idx < len(current) else None
        else:
            current = current.get(part) if isinstance(current, dict) else None
        if current is None:
            return None
    return current


def first_source_ref(field_meta: dict[str, Any] | None) -> dict[str, Any] | None:
    if not field_meta:
        return None
    refs = field_meta.get("source_refs")
    if isinstance(refs, list) and refs:
        return refs[0]
    for alt in field_meta.get("alternatives") or []:
        refs = alt.get("source_refs") if isinstance(alt, dict) else None
        if isinstance(refs, list) and refs:
            ref = dict(refs[0])
            ref["_alternative_value"] = alt.get("value")
            return ref
    return None


def render_placeholder(title: str, lines: list[str]) -> str:
    width, height = 1100, 660
    image = Image.new("RGB", (width, height), "#f7f7f5")
    draw = ImageDraw.Draw(image)
    try:
        title_font = ImageFont.truetype("/System/Library/Fonts/ヒラギノ角ゴシック W6.ttc", 34)
        body_font = ImageFont.truetype("/System/Library/Fonts/ヒラギノ角ゴシック W3.ttc", 25)
    except Exception:
        title_font = ImageFont.load_default()
        body_font = ImageFont.load_default()
    draw.rectangle((0, 0, width - 1, height - 1), outline="#d7d7d2", width=3)
    draw.text((42, 42), title, fill="#1f2937", font=title_font)
    y = 118
    for line in lines:
        for wrapped in wrap_visual(line, 44):
            draw.text((42, y), wrapped, fill="#374151", font=body_font)
            y += 40
        y += 10
    return image_to_data_uri(image)


def wrap_visual(text: str, width: int) -> list[str]:
    chunks = []
    current = ""
    for char in text:
        current += char
        visual_len = sum(2 if ord(c) > 127 else 1 for c in current)
        if visual_len >= width:
            chunks.append(current)
            current = ""
    if current:
        chunks.append(current)
    return chunks or [""]


def image_to_data_uri(image: Image.Image, max_width: int = 1300) -> str:
    if image.width > max_width:
        ratio = max_width / image.width
        image = image.resize((max_width, max(1, int(image.height * ratio))), Image.Resampling.LANCZOS)
    if image.mode != "RGB":
        image = image.convert("RGB")
    buf = io.BytesIO()
    image.save(buf, format="JPEG", quality=76, optimize=True)
    payload = base64.b64encode(buf.getvalue()).decode("ascii")
    return f"data:image/jpeg;base64,{payload}"


def render_pdf_page(pdf_path: Path, page_number: int, quote: str | None = None) -> tuple[str, str]:
    try:
        doc = fitz.open(pdf_path)
        page_index = min(max(page_number - 1, 0), len(doc) - 1)
        page = doc[page_index]
        clip = None
        if quote:
            for term in search_terms(quote):
                rects = page.search_for(term)
                if rects:
                    clip = rects[0]
                    clip.x0 = max(0, clip.x0 - 90)
                    clip.y0 = max(0, clip.y0 - 80)
                    clip.x1 = min(page.rect.x1, clip.x1 + 400)
                    clip.y1 = min(page.rect.y1, clip.y1 + 110)
                    break
        if clip:
            pix = page.get_pixmap(matrix=fitz.Matrix(3.0, 3.0), clip=clip, alpha=False)
            scope = f"p.{page_index + 1} crop"
        else:
            pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
            scope = f"p.{page_index + 1} full page"
        image = Image.open(io.BytesIO(pix.tobytes("png")))
        doc.close()
        return image_to_data_uri(image), scope
    except Exception as exc:
        return render_placeholder("PDF rendering failed", [str(pdf_path), str(exc)]), "render failed"


def search_terms(value: Any) -> list[str]:
    text = str(value or "").strip()
    terms = [text]
    if "　" in text:
        terms.append(text.replace("　", " "))
    if "," in text:
        terms.append(text.replace(",", " "))
    digits = re.sub(r"\D", "", text)
    if len(digits) >= 6:
        terms.append(digits)
    return [t for t in dict.fromkeys(terms) if t]


def render_rasens_evidence(case_dir: Path, fields: list[FieldMismatch]) -> EvidenceImage:
    docs = build_document_map(case_dir)
    rasens = docs.get("rasens_application_output")
    if not rasens:
        return EvidenceImage(
            "1. RASENSに書かれている内容",
            "missing",
            render_placeholder("RASENS output not found", ["output_manifestにrasens_application_outputがありません。"]),
            "RASENS output file missing",
        )
    pdf_path = Path(rasens["abs_path"])
    field = first_searchable_field(fields)
    page = find_pdf_page_for_terms(pdf_path, [field.expected, field.generated]) or RASENS_PAGE_HINTS.get(field.path, 1)
    quote = field.expected if field.expected not in {"", "False", "True"} else field.generated
    image_uri, scope = render_pdf_page(pdf_path, page, quote)
    return EvidenceImage(
        "1. RASENSに書かれている内容",
        "RASENS",
        image_uri,
        f"{pdf_path.name} / {scope}",
    )


def first_searchable_field(fields: list[FieldMismatch]) -> FieldMismatch:
    for field in fields:
        if field.expected not in {"", "False", "True"}:
            return field
    for field in fields:
        if field.generated not in {"", "False", "True"}:
            return field
    return fields[0]


def find_pdf_page_for_terms(pdf_path: Path, values: list[Any]) -> int | None:
    try:
        doc = fitz.open(pdf_path)
        for page_index, page in enumerate(doc):
            text = page.get_text("text")
            haystack = normalize_text(text)
            for value in values:
                for term in search_terms(value):
                    n = normalize_text(term)
                    if n and n in haystack:
                        doc.close()
                        return page_index + 1
        doc.close()
    except Exception:
        return None
    return None


def render_source_ref_evidence(
    ref: dict[str, Any] | None,
    docs: dict[str, dict[str, Any]],
    title: str,
    fallback_lines: list[str],
) -> EvidenceImage:
    if not ref:
        return EvidenceImage(title, "no source_ref", render_placeholder("根拠なし", fallback_lines), "source_refsなし")
    document_id = ref.get("document_id")
    doc = docs.get(document_id)
    quote = ref.get("text_quote") or ref.get("_alternative_value") or ""
    if not doc:
        return EvidenceImage(
            title,
            "missing document",
            render_placeholder("document not found", [f"document_id: {document_id}", str(quote)]),
            f"document_id={document_id}",
        )
    path = Path(doc["abs_path"])
    ext = doc.get("extension", path.suffix.lstrip(".")).lower()
    if ext == "pdf":
        page = int(ref.get("page") or find_pdf_page_for_terms(path, [quote]) or 1)
        image_uri, scope = render_pdf_page(path, page, str(quote))
        return EvidenceImage(title, "AI input", image_uri, f"{doc.get('document_role', document_id)} / {scope} / quote: {quote}")
    snippet = text_snippet_for_document(path, ext, str(quote))
    image_uri = render_text_snippet(path.name, snippet, str(quote))
    return EvidenceImage(title, "AI input", image_uri, f"{doc.get('document_role', document_id)} / quote: {quote}")


def text_snippet_for_document(path: Path, ext: str, quote: str) -> list[str]:
    lines = extract_text_lines(path, ext)
    needle = normalize_text(quote)
    if needle:
        for i, line in enumerate(lines):
            if needle in normalize_text(line):
                start = max(0, i - 3)
                end = min(len(lines), i + 4)
                return lines[start:end]
    return lines[:8] if lines else ["テキスト抽出できませんでした。"]


def extract_text_lines(path: Path, ext: str) -> list[str]:
    try:
        if ext == "docx":
            from docx import Document

            doc = Document(str(path))
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
            return lines
        if ext == "xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(path, data_only=True, read_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    cells = [str(v).strip() for v in row if v not in (None, "")]
                    if cells:
                        lines.append(f"{ws.title} row {idx}: " + " | ".join(cells))
            return lines
        return [path.read_text(encoding="utf-8", errors="ignore")]
    except Exception as exc:
        return [f"抽出エラー: {exc}"]


def render_text_snippet(file_name: str, lines: list[str], quote: str) -> str:
    width, height = 1300, 760
    image = Image.new("RGB", (width, height), "#ffffff")
    draw = ImageDraw.Draw(image)
    try:
        title_font = ImageFont.truetype("/System/Library/Fonts/ヒラギノ角ゴシック W6.ttc", 28)
        body_font = ImageFont.truetype("/System/Library/Fonts/ヒラギノ角ゴシック W3.ttc", 23)
    except Exception:
        title_font = ImageFont.load_default()
        body_font = ImageFont.load_default()
    draw.rectangle((0, 0, width - 1, height - 1), outline="#d1d5db", width=2)
    draw.text((34, 30), file_name, fill="#111827", font=title_font)
    y = 92
    normalized_quote = normalize_text(quote)
    for raw_line in lines[:10]:
        marker = ">"
        line = raw_line
        is_hit = normalized_quote and normalized_quote in normalize_text(line)
        bg = "#fff7ed" if is_hit else "#ffffff"
        if is_hit:
            draw.rectangle((24, y - 8, width - 24, y + 64), fill=bg)
        for wrapped in wrap_visual(f"{marker} {line}", 96):
            draw.text((34, y), wrapped, fill="#1f2937", font=body_font)
            y += 34
            if y > height - 42:
                break
        y += 10
        if y > height - 42:
            break
    return image_to_data_uri(image)


def render_correct_source_evidence(
    case_dir: Path,
    docs: dict[str, dict[str, Any]],
    fields: list[FieldMismatch],
    field_meta: dict[str, dict[str, Any]],
) -> EvidenceImage:
    searchable = [f.expected for f in fields if f.expected not in {"", "False", "True"}]
    if searchable:
        found = find_input_document_evidence(docs, searchable)
        if found:
            return found
    related_ref = related_source_ref(fields, field_meta)
    if related_ref:
        lines = [
            "正解値そのものではなく、判断対象になった入力箇所です。",
            "空欄/Falseの正解は、この箇所を役職・13桁法人番号として採用しない判断です。",
        ]
        return render_source_ref_evidence(related_ref, docs, "4. 実際の正解値の入力資料", lines)
    generated_terms = [f.generated for f in fields if f.generated not in {"", "False", "True"}]
    if generated_terms:
        found = find_input_document_evidence(docs, generated_terms)
        if found:
            found.title = "4. 実際の正解値の入力資料"
            found.badge = "judgement source"
            found.caption += " / 正解はこの箇所を採用しない判断"
            return found
    return EvidenceImage(
        "4. 実際の正解値の入力資料",
        "not located",
        render_placeholder(
            "正解根拠を自動特定できません",
            [
                "expected値を入力資料内で自動検索しましたが、該当箇所を特定できませんでした。",
                "HTML内の値テーブルとRASENS画像で確認してください。",
            ],
        ),
        "input evidence not located automatically",
    )


def find_input_document_evidence(docs: dict[str, dict[str, Any]], values: list[str]) -> EvidenceImage | None:
    input_docs = [d for d in docs.values() if d.get("use_as_input") is True]
    for doc in input_docs:
        path = Path(doc["abs_path"])
        ext = doc.get("extension", path.suffix.lstrip(".")).lower()
        if ext == "pdf":
            page = find_pdf_page_for_terms(path, values)
            if page:
                quote = next((v for v in values if v), "")
                image_uri, scope = render_pdf_page(path, page, quote)
                return EvidenceImage(
                    "4. 実際の正解値の入力資料",
                    "found",
                    image_uri,
                    f"{doc.get('document_role', doc.get('document_id'))} / {scope}",
                )
        elif ext in {"docx", "xlsx"}:
            lines = extract_text_lines(path, ext)
            for value in values:
                needle = normalize_text(value)
                if needle and any(needle in normalize_text(line) for line in lines):
                    snippet = text_snippet_for_document(path, ext, value)
                    return EvidenceImage(
                        "4. 実際の正解値の入力資料",
                        "found",
                        render_text_snippet(path.name, snippet, value),
                        f"{doc.get('document_role', doc.get('document_id'))} / matched expected value",
                    )
    return None


def related_source_ref(fields: list[FieldMismatch], field_meta: dict[str, dict[str, Any]]) -> dict[str, Any] | None:
    sibling_order = [
        "employer.corporate_number",
        "employer.has_corporate_number",
        "employment.position_title",
        "employment.has_position",
    ]
    field_paths = {f.path for f in fields}
    for path in sibling_order:
        if path in field_paths:
            ref = first_source_ref(field_meta.get(path))
            if ref:
                return ref
    for field in fields:
        ref = first_source_ref(field_meta.get(field.path))
        if ref:
            return ref
    return None


def classify_group(fields: list[FieldMismatch]) -> tuple[str, str, str, str]:
    paths = {f.path for f in fields}
    if paths & CORPORATE_PATHS:
        return (
            "corporate",
            "法人番号・有無の不一致",
            "13桁の法人番号と12桁の会社法人等番号の区別が原因です。AIは資料上の番号を根拠として拾っていますが、アプリ上の正解は13桁法人番号が確認できない場合は空欄/Falseです。",
            "法人番号は「13桁だけ採用、12桁は採用しない」を抽出指示と後処理の両方で明示し、has_corporate_numberはcorporate_numberの非空と連動させます。",
        )
    if paths & POSITION_PATHS:
        return (
            "position",
            "役職有無・役職名の不一致",
            "職種や業務内容を役職名として扱ったことが原因です。役職欄は正式な肩書きがある場合だけ埋めるべきで、業務説明とは分ける必要があります。",
            "position_titleは正式役職のみ、職種・担当業務は別fieldへ、has_positionはposition_titleの非空と連動させます。",
        )
    return (
        "address",
        f"{FIELD_LABELS.get(fields[0].path, fields[0].path)}の不一致",
        "住所・出生地・本国住所の粒度と参照元の優先順位がずれています。AIは入力資料内の別住所や既存申請添付の住所を拾い、アプリで最終的に持つべき統一形式と一致していません。",
        "出生地、本国住所、所属機関住所の意味を短く固定し、提出済み申請書より一次資料・会社資料を優先するルールを明示します。",
    )


def group_mismatches(run_dir: Path) -> list[tuple[Path, list[FieldMismatch]]]:
    groups: list[tuple[Path, list[FieldMismatch]]] = []
    for generated_dir in sorted([p for p in run_dir.iterdir() if p.is_dir() and (p / "comparison_report.md").exists()]):
        rows = parse_detail_rows(generated_dir / "comparison_report.md")
        mismatches: list[FieldMismatch] = []
        for row in rows:
            path = f"{row.get('大項目')}.{row.get('小項目')}"
            if row.get("判定") != "❌ 不一致" or path not in TARGET_PATHS:
                continue
            if path in ADDRESS_PATHS:
                category = "address"
            elif path in CORPORATE_PATHS:
                category = "corporate"
            else:
                category = "position"
            mismatches.append(
                FieldMismatch(
                    path=path,
                    expected=row.get("正解データ", ""),
                    generated=row.get("AI出力", ""),
                    category=category,
                )
            )
        by_key: dict[str, list[FieldMismatch]] = {}
        for item in mismatches:
            if item.path in ADDRESS_PATHS:
                key = item.path
            elif item.path in CORPORATE_PATHS:
                key = "employer.corporate"
            else:
                key = "employment.position"
            by_key.setdefault(key, []).append(item)
        for items in by_key.values():
            groups.append((generated_dir, sorted(items, key=lambda f: f.path)))
    return groups


def build_cards(run_dir: Path) -> list[Card]:
    grouped = group_mismatches(run_dir)
    cards: list[Card] = []
    for idx, (generated_dir, fields) in enumerate(grouped, start=1):
        case_dir = find_fixture_dir(generated_dir.name)
        docs = build_document_map(case_dir)
        field_meta = load_json(generated_dir / "field_metadata.json")
        category, issue_title, cause, improvement = classify_group(fields)
        ref = related_source_ref(fields, field_meta)
        rasens = render_rasens_evidence(case_dir, fields)
        if ref:
            ai_source = render_source_ref_evidence(
                ref,
                docs,
                "3. AIが推測した値を取った箇所",
                ["AI出力にsource_refsがありません。値テーブルとreview.jsonを確認してください。"],
            )
        else:
            generated_terms = [f.generated for f in fields if f.generated not in {"", "False", "True"}]
            fallback = find_input_document_evidence(docs, generated_terms) if generated_terms else None
            if fallback:
                fallback.title = "3. AIが推測した値を取った箇所"
                fallback.badge = "value search"
                fallback.caption += " / source_refsなしのためAI出力値で逆検索"
                ai_source = fallback
            else:
                ai_source = render_source_ref_evidence(
                    None,
                    docs,
                    "3. AIが推測した値を取った箇所",
                    ["AI出力にsource_refsがなく、AI出力値の入力資料内検索でも特定できませんでした。"],
                )
        correct_source = render_correct_source_evidence(case_dir, docs, fields, field_meta)
        fixture_label = f"F{fixture_index(run_dir, generated_dir)}"
        cards.append(
            Card(
                fixture_label=fixture_label,
                case_dir=case_dir,
                generated_dir=generated_dir,
                fields=fields,
                category=category,
                issue_title=issue_title,
                cause=cause,
                improvement=improvement,
                rasens=rasens,
                ai_source=ai_source,
                correct_source=correct_source,
            )
        )
    return cards


def fixture_index(run_dir: Path, generated_dir: Path) -> int:
    dirs = sorted([p for p in run_dir.iterdir() if p.is_dir() and (p / "comparison_report.md").exists()])
    return dirs.index(generated_dir) + 1


def counts(cards: list[Card]) -> dict[str, int]:
    result = {"fields": 0, "cards": len(cards), "address": 0, "corporate": 0, "position": 0}
    for card in cards:
        result["fields"] += len(card.fields)
        for field in card.fields:
            result[field.category] += 1
    return result


def escape(value: Any) -> str:
    return html.escape(str(value if value is not None else ""), quote=True)


def render_html(cards: list[Card], run_dir: Path) -> str:
    c = counts(cards)
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    card_html = "\n".join(render_card(card, i) for i, card in enumerate(cards, start=1))
    fixtures = sorted({card.fixture_label for card in cards}, key=lambda x: int(x[1:]))
    fixture_options = "\n".join(f'<option value="{f}">{f}</option>' for f in fixtures)
    return f"""<!doctype html>
<html lang="ja">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>住所・法人番号・役職の不一致 証拠一覧</title>
<style>
:root {{ --bg:#f6f7f9; --panel:#fff; --text:#111827; --muted:#6b7280; --line:#d9dde5; --accent:#2563eb; --warn:#9a3412; }}
* {{ box-sizing: border-box; }}
body {{ margin:0; background:var(--bg); color:var(--text); font:14px/1.55 -apple-system,BlinkMacSystemFont,"Hiragino Sans","Yu Gothic",Meiryo,sans-serif; }}
header {{ position:sticky; top:0; z-index:20; background:rgba(246,247,249,.96); border-bottom:1px solid var(--line); backdrop-filter: blur(8px); }}
.wrap {{ max-width:1500px; margin:0 auto; padding:18px 22px; }}
h1 {{ margin:0 0 8px; font-size:24px; letter-spacing:0; }}
.sub {{ color:var(--muted); }}
.summary {{ display:grid; grid-template-columns:repeat(5,minmax(130px,1fr)); gap:10px; margin-top:14px; }}
.metric {{ background:var(--panel); border:1px solid var(--line); border-radius:8px; padding:10px 12px; }}
.metric b {{ display:block; font-size:22px; }}
.controls {{ display:flex; flex-wrap:wrap; gap:10px; align-items:center; margin-top:14px; }}
input, select, button {{ border:1px solid var(--line); background:#fff; border-radius:7px; padding:9px 10px; font:inherit; }}
input {{ flex:1; min-width:260px; }}
button {{ cursor:pointer; }}
.note {{ margin-top:12px; color:#4b5563; background:#fff7ed; border:1px solid #fed7aa; border-radius:8px; padding:10px 12px; }}
main.wrap {{ padding-top:18px; }}
.card {{ background:var(--panel); border:1px solid var(--line); border-radius:8px; margin:0 0 16px; overflow:hidden; }}
.card summary {{ cursor:pointer; padding:14px 16px; display:flex; gap:12px; align-items:center; justify-content:space-between; }}
.title {{ display:flex; gap:10px; align-items:center; flex-wrap:wrap; }}
.badge {{ display:inline-flex; align-items:center; border:1px solid var(--line); border-radius:999px; padding:2px 8px; color:#374151; background:#f9fafb; font-size:12px; }}
.badge.address {{ border-color:#bfdbfe; background:#eff6ff; }}
.badge.corporate {{ border-color:#fde68a; background:#fffbeb; }}
.badge.position {{ border-color:#ddd6fe; background:#f5f3ff; }}
.content {{ padding:0 16px 16px; }}
table {{ width:100%; border-collapse:collapse; margin:0 0 14px; table-layout:fixed; }}
th,td {{ border:1px solid var(--line); padding:8px 9px; vertical-align:top; word-break:break-word; }}
th {{ background:#f9fafb; text-align:left; color:#374151; }}
.explain {{ display:grid; grid-template-columns:1fr 1fr; gap:12px; margin:12px 0 14px; }}
.explain > div {{ border:1px solid var(--line); border-radius:8px; padding:10px 12px; background:#fcfcfd; }}
.evidence {{ display:grid; grid-template-columns:repeat(3,minmax(0,1fr)); gap:12px; }}
.shot {{ border:1px solid var(--line); border-radius:8px; overflow:hidden; background:#fff; }}
.shot h3 {{ margin:0; padding:10px 12px 0; font-size:14px; }}
.shot .meta {{ color:var(--muted); padding:2px 12px 9px; font-size:12px; }}
.shot img {{ display:block; width:100%; height:320px; object-fit:contain; background:#f3f4f6; cursor:zoom-in; border-top:1px solid var(--line); }}
.hidden {{ display:none !important; }}
#modal {{ position:fixed; inset:0; background:rgba(17,24,39,.86); display:none; align-items:center; justify-content:center; z-index:100; padding:24px; }}
#modal.open {{ display:flex; }}
#modal img {{ max-width:96vw; max-height:92vh; background:#fff; }}
#modal button {{ position:fixed; top:16px; right:16px; }}
@media (max-width: 1050px) {{ .summary {{ grid-template-columns:repeat(2,1fr); }} .evidence,.explain {{ grid-template-columns:1fr; }} .shot img {{ height:260px; }} }}
@media print {{ header {{ position:static; }} .controls {{ display:none; }} .card {{ break-inside:avoid; }} .shot img {{ height:auto; }} }}
</style>
</head>
<body>
<header>
  <div class="wrap">
    <h1>住所・法人番号・役職の不一致 証拠一覧</h1>
    <div class="sub">run: {escape(run_dir.name)} / generated: {escape(generated_at)} / restricted local report</div>
    <div class="summary">
      <div class="metric"><span>field不一致</span><b>{c["fields"]}</b></div>
      <div class="metric"><span>カード</span><b>{c["cards"]}</b></div>
      <div class="metric"><span>住所・出生地・本国住所</span><b>{c["address"]}</b></div>
      <div class="metric"><span>法人番号・有無</span><b>{c["corporate"]}</b></div>
      <div class="metric"><span>役職有無・役職名</span><b>{c["position"]}</b></div>
    </div>
    <div class="controls">
      <input id="q" type="search" placeholder="検索: field名、値、原因、改善策">
      <select id="category"><option value="">全分類</option><option value="address">住所</option><option value="corporate">法人番号</option><option value="position">役職</option></select>
      <select id="fixture"><option value="">全fixture</option>{fixture_options}</select>
      <button id="expand">全展開</button><button id="collapse">全折りたたみ</button>
    </div>
    <div class="note">このHTMLは実PIIを含みます。git管理外のeval_runs配下でローカル確認用途に限定してください。</div>
  </div>
</header>
<main class="wrap" id="cards">
{card_html}
</main>
<div id="modal"><button>閉じる</button><img alt=""></div>
<script>
const cards = [...document.querySelectorAll('.card')];
const q = document.getElementById('q');
const category = document.getElementById('category');
const fixture = document.getElementById('fixture');
function applyFilter() {{
  const text = q.value.trim().toLowerCase();
  for (const card of cards) {{
    const okText = !text || card.innerText.toLowerCase().includes(text);
    const okCat = !category.value || card.dataset.category === category.value;
    const okFix = !fixture.value || card.dataset.fixture === fixture.value;
    card.classList.toggle('hidden', !(okText && okCat && okFix));
  }}
}}
[q, category, fixture].forEach(el => el.addEventListener('input', applyFilter));
document.getElementById('expand').addEventListener('click', () => cards.forEach(c => c.open = true));
document.getElementById('collapse').addEventListener('click', () => cards.forEach(c => c.open = false));
const modal = document.getElementById('modal');
document.querySelectorAll('.shot img').forEach(img => img.addEventListener('click', () => {{
  modal.querySelector('img').src = img.src;
  modal.classList.add('open');
}}));
modal.addEventListener('click', () => modal.classList.remove('open'));
modal.querySelector('button').addEventListener('click', () => modal.classList.remove('open'));
</script>
</body>
</html>
"""


def render_card(card: Card, index: int) -> str:
    field_rows = "\n".join(
        f"<tr><td>{escape(FIELD_LABELS.get(f.path, f.path))}</td><td><code>{escape(f.path)}</code></td><td>{escape(f.expected)}</td><td>{escape(f.generated)}</td></tr>"
        for f in card.fields
    )
    shots = "\n".join(render_shot(s) for s in [card.rasens, card.ai_source, card.correct_source])
    search_blob = " ".join([card.issue_title, card.cause, card.improvement] + [f.path + " " + f.expected + " " + f.generated for f in card.fields])
    return f"""
<details class="card" data-category="{escape(card.category)}" data-fixture="{escape(card.fixture_label)}" data-search="{escape(search_blob)}" open>
  <summary>
    <div class="title"><span class="badge">{index:02d}</span><span class="badge">{escape(card.fixture_label)}</span><span class="badge {escape(card.category)}">{escape(category_label(card.category))}</span><strong>{escape(card.issue_title)}</strong></div>
    <span class="sub">{len(card.fields)} field</span>
  </summary>
  <div class="content">
    <table>
      <thead><tr><th style="width:16%">項目</th><th style="width:24%">path</th><th>正解データ</th><th>AIが推測した値</th></tr></thead>
      <tbody>{field_rows}</tbody>
    </table>
    <div class="explain">
      <div><strong>原因</strong><br>{escape(card.cause)}</div>
      <div><strong>改善策</strong><br>{escape(card.improvement)}</div>
    </div>
    <div class="evidence">{shots}</div>
  </div>
</details>
"""


def render_shot(evidence: EvidenceImage) -> str:
    return f"""
<section class="shot">
  <h3>{escape(evidence.title)} <span class="badge">{escape(evidence.badge)}</span></h3>
  <div class="meta">{escape(evidence.caption)}</div>
  <img src="{evidence.image_uri}" alt="{escape(evidence.title)}">
</section>
"""


def category_label(category: str) -> str:
    return {"address": "住所", "corporate": "法人番号", "position": "役職"}.get(category, category)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    run_dir = (ROOT / args.run_dir).resolve() if not args.run_dir.is_absolute() else args.run_dir
    output = (ROOT / args.output).resolve() if not args.output.is_absolute() else args.output
    cards = build_cards(run_dir)
    c = counts(cards)
    if c["fields"] != 39 or c["address"] != 19 or c["corporate"] != 10 or c["position"] != 10 or c["cards"] != 31:
        raise SystemExit(f"unexpected target counts: {c}")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(render_html(cards, run_dir), encoding="utf-8")
    print(json.dumps({"output": str(output), **c}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
